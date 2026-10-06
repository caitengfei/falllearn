# -*- coding: utf-8 -*-
"""教师端知识库管理（P2.15）：knowledge/ 目录为唯一事实源（学生端 /kb 与 AI 问答同源）。
文档列表 + 分类统计 + 手动新增/编辑 + 文件上传（txt/pdf/docx 提取文字）+ 回收站式删除。
预留向量/重排模型 API 接入（OpenAI 兼容）：不配置 = 纯本地行为不变；
配置后教师可分块向量化，学生端检索自动融合（见 kb.py _vector_enhance）。"""
import io
import json
import os
import re
import struct
import time
import urllib.request

from fastapi import APIRouter, Depends, File, Form, HTTPException, UploadFile
from pydantic import BaseModel

from . import db, llm_direct
from .auth import require_teacher

router = APIRouter(prefix="/api/admin/kb", tags=["kb_admin"])

DIMS = ["01-岗", "02-课", "03-赛", "04-证", "05-元数据"]
MAX_FILE_MB = 8            # 上传大小上限（教学文档）
MAX_TEXT_CHARS = 400_000   # 单篇正文上限（防 AI 上下文/检索膨胀）
CHUNK_SIZE = 500           # 向量化分块（字符）
KB_DIR = llm_direct.KB_DIR


def _invalidate():
    """写后失效 llm_direct 懒加载缓存：学生端 /kb 与 AI 问答立即可见（不等 5 分钟 TTL）。"""
    llm_direct._kb_cache["at"] = 0


def _check_path(dim: str, title: str):
    dim = (dim or "").strip()
    if dim not in DIMS:
        raise HTTPException(400, "分类必须是：" + " / ".join(DIMS))
    title = re.sub(r'[\\/:*?"<>|\r\n\t]', "", (title or "").strip())[:80]
    if not title:
        raise HTTPException(400, "请填写标题")
    return dim, title


def _full(dim: str, title: str) -> str:
    p = os.path.join(KB_DIR, dim, title + ".md")
    base = os.path.realpath(KB_DIR)
    real = os.path.realpath(p)
    try:
        if os.path.commonpath([base, real]) != base:
            raise HTTPException(400, "非法路径")
    except ValueError:
        raise HTTPException(400, "非法路径")
    return real


def _rel(dim: str, title: str) -> str:
    return f"{dim}/{title}.md"


def _snippet(text: str) -> str:
    for line in text.splitlines():
        line = line.strip()
        if line and not line.startswith(("#", ">", "|", "-")):
            return line[:80]
    return ""


def _trash(rel: str) -> bool:
    """删除 = 移入 knowledge/.trash/（可人工恢复），并清理该文档向量。"""
    src = os.path.join(KB_DIR, *rel.split("/"))
    if not os.path.isfile(src):
        return False
    tdir = os.path.join(KB_DIR, ".trash")
    os.makedirs(tdir, exist_ok=True)
    os.replace(src, os.path.join(tdir, time.strftime("%Y%m%d%H%M%S_") + os.path.basename(src)))
    d = db.get_db()
    d.execute("DELETE FROM kb_embeddings WHERE doc_path=?", (rel,))
    d.commit()
    d.close()
    return True


def _clean_vecs(rel: str):
    d = db.get_db()
    d.execute("DELETE FROM kb_embeddings WHERE doc_path=?", (rel,))
    d.commit()
    d.close()


# ---------------- 列表 / 统计 ----------------

@router.get("")
def kb_list(u: dict = Depends(require_teacher)):
    """文档列表 + 分类统计（各维度篇数/字数）+ 维度清单。"""
    items, stats = [], {d: {"count": 0, "chars": 0} for d in DIMS}
    stats["其他"] = {"count": 0, "chars": 0}
    for rel, text in llm_direct.kb_fresh(0):  # max_age=0 强制刷新（管理端低频）
        top = rel.split("/", 1)[0]
        d = top if top in DIMS else "其他"
        stats[d]["count"] += 1
        stats[d]["chars"] += len(text)
        title = rel.rsplit("/", 1)[-1].rsplit(".md", 1)[0]
        full = os.path.join(KB_DIR, *rel.split("/"))
        mt = int(os.path.getmtime(full)) if os.path.isfile(full) else 0
        items.append({"path": rel, "title": title, "dim": d, "chars": len(text),
                      "mtime": mt, "snippet": _snippet(text)})
    items.sort(key=lambda x: x["path"])
    return {"items": items, "stats": stats, "total": len(items), "dims": DIMS}


# ---------------- 新增 / 编辑 ----------------

class KbIn(BaseModel):
    dim: str
    title: str
    content: str
    path: str = ""  # 非空 = 编辑既有文档（改名时旧文件移回收站）


@router.post("")
def kb_save(b: KbIn, u: dict = Depends(require_teacher)):
    dim, title = _check_path(b.dim, b.title)
    text = (b.content or "").strip()
    if len(text) < 20:
        raise HTTPException(400, "正文过短（至少 20 字）")
    if len(text) > MAX_TEXT_CHARS:
        raise HTTPException(400, f"正文过长（上限 {MAX_TEXT_CHARS // 1000} 千字）")
    # 编辑改名：旧文件移回收站（按原路径）
    old = (b.path or "").replace("\\", "/").lstrip("/")
    rel = _rel(dim, title)
    is_edit = bool(old)
    if old and old != rel:
        parts = old.split("/")
        if len(parts) == 2 and old.endswith(".md"):
            _trash(old)
    full = _full(dim, title)
    os.makedirs(os.path.dirname(full), exist_ok=True)
    with open(full, "w", encoding="utf-8") as fh:
        fh.write(f"# {title}\n\n" + text + "\n")
    _clean_vecs(rel)
    _invalidate()
    return {"ok": True, "path": rel, "created": not is_edit}


# ---------------- 文件上传（txt / pdf / docx / md） ----------------

def _extract_text(name: str, raw: bytes) -> str:
    low = name.lower()
    if low.endswith(".txt"):
        for enc in ("utf-8", "gbk"):
            try:
                return raw.decode(enc)
            except UnicodeDecodeError:
                continue
        raise HTTPException(400, "txt 编码无法识别（支持 UTF-8 / GBK）")
    if low.endswith(".md"):
        return raw.decode("utf-8", errors="ignore")
    if low.endswith(".pdf"):
        try:
            from pypdf import PdfReader
        except ImportError:
            raise HTTPException(400, "服务器缺少 pypdf 组件（pip install pypdf 后重试）")
        r = PdfReader(io.BytesIO(raw))
        return "\n\n".join((pg.extract_text() or "") for pg in r.pages)
    if low.endswith(".docx"):
        try:
            import docx
        except ImportError:
            raise HTTPException(400, "服务器缺少 python-docx 组件（pip install python-docx 后重试）")
        doc = docx.Document(io.BytesIO(raw))
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for t in doc.tables:
            for row in t.rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if cells:
                    parts.append(" ".join(cells))
        return "\n".join(parts)
    raise HTTPException(400, "仅支持 .txt / .pdf / .docx / .md")


@router.post("/upload")
async def kb_upload(file: UploadFile = File(...), dim: str = Form(""),
                    title: str = Form(""), u: dict = Depends(require_teacher)):
    raw = await file.read()
    if len(raw) > MAX_FILE_MB * 1024 * 1024:
        raise HTTPException(400, f"文件过大（上限 {MAX_FILE_MB}MB）")
    name = (file.filename or "untitled").strip()
    if not re.search(r"\.(txt|pdf|docx|md)$", name, re.I):
        raise HTTPException(400, "请上传 .txt / .pdf / .docx / .md 文件")
    text = _extract_text(name, raw).strip()
    if len(text) < 20:
        raise HTTPException(400, "未能从文件中提取到有效文字（建议手动录入）")
    if len(text) > MAX_TEXT_CHARS:
        text = text[:MAX_TEXT_CHARS]
    # 原件存档（可追溯；本地目录，不出服务器）
    updir = os.path.join(KB_DIR, "uploads")
    os.makedirs(updir, exist_ok=True)
    safe = re.sub(r"[^\w.\-]", "_", name)
    oname = time.strftime("%Y%m%d%H%M%S_") + safe
    with open(os.path.join(updir, oname), "wb") as fh:
        fh.write(raw)
    dim, title = _check_path(dim, (title or os.path.splitext(name)[0]).strip())
    full = _full(dim, title)
    os.makedirs(os.path.dirname(full), exist_ok=True)
    with open(full, "w", encoding="utf-8") as fh:
        fh.write(f"# {title}\n\n> 来源：教师上传 `{name}`（{time.strftime('%Y-%m-%d')}）\n\n" + text + "\n")
    _clean_vecs(_rel(dim, title))
    _invalidate()
    return {"ok": True, "path": _rel(dim, title), "chars": len(text)}


# ---------------- 删除（回收站） ----------------

@router.delete("")
def kb_delete(path: str, u: dict = Depends(require_teacher)):
    rel = (path or "").replace("\\", "/").lstrip("/")
    parts = rel.split("/")
    if len(parts) != 2 or not rel.endswith(".md") or ".." in parts:
        raise HTTPException(400, "非法路径")
    if not _trash(rel):
        raise HTTPException(404, "文档不存在")
    _invalidate()
    return {"ok": True}


# ---------------- 模型配置（向量 / 重排；settings 表） ----------------

def _cfg_view(c: dict) -> dict:
    out = {k: c.get(k, "") for k in ("base_url", "api_key", "model")}
    out["configured"] = bool(out["base_url"] and out["api_key"] and out["model"])
    k = out.get("api_key") or ""
    out["api_key"] = (k[:6] + "****" + k[-4:]) if len(k) > 12 else ("****" if k else "")
    return out


@router.get("/config")
def kb_cfg(u: dict = Depends(require_teacher)):
    d = db.get_db()
    rows = {r["key"]: r["value"] for r in d.execute(
        "SELECT key, value FROM settings WHERE key IN ('kb_embed','kb_rerank')")}
    d.close()
    def _v(key):
        try:
            return _cfg_view(json.loads(rows.get(key) or "{}"))
        except json.JSONDecodeError:
            return _cfg_view({})
    return {"embed": _v("kb_embed"), "rerank": _v("kb_rerank")}


class CfgIn(BaseModel):
    base_url: str = ""
    api_key: str = ""
    model: str = ""


@router.put("/config/{kind}")
def kb_cfg_save(kind: str, b: CfgIn, u: dict = Depends(require_teacher)):
    if kind not in ("embed", "rerank"):
        raise HTTPException(400, "非法类型")
    key = "kb_embed" if kind == "embed" else "kb_rerank"
    d = db.get_db()
    old = d.execute("SELECT value FROM settings WHERE key=?", (key,)).fetchone()
    try:
        oldv = json.loads(old["value"] if old and old["value"] else "{}")
    except json.JSONDecodeError:
        oldv = {}
    # 前端回传的掩码值（含 ****）= 未修改，保留原密钥
    if b.api_key and "****" in b.api_key:
        b.api_key = oldv.get("api_key", "")
    v = {"base_url": (b.base_url or "").strip().rstrip("/"),
         "api_key": (b.api_key or "").strip(),
         "model": (b.model or "").strip()}
    d.execute("INSERT INTO settings(key,value) VALUES(?,?) "
              "ON CONFLICT(key) DO UPDATE SET value=excluded.value",
              (key, json.dumps(v, ensure_ascii=False)))
    d.commit()
    d.close()
    return {"ok": True}


def _load_cfg(kind: str):
    d = db.get_db()
    r = d.execute("SELECT value FROM settings WHERE key=?",
                  ("kb_embed" if kind == "embed" else "kb_rerank",)).fetchone()
    d.close()
    try:
        c = json.loads(r["value"] if r and r["value"] else "{}")
    except json.JSONDecodeError:
        c = {}
    return c if (c.get("base_url") and c.get("api_key") and c.get("model")) else None


# ---------------- 向量化（OpenAI 兼容 /embeddings） ----------------

def _post_json(url: str, headers: dict, payload: dict, timeout: int = 90):
    req = urllib.request.Request(
        url, data=json.dumps(payload).encode("utf-8"),
        headers={"content-type": "application/json", **headers}, method="POST")
    with urllib.request.urlopen(req, timeout=timeout) as resp:
        return json.loads(resp.read().decode("utf-8"))


def _to_blob(vec) -> bytes:
    return struct.pack(f"<{len(vec)}f", *vec)


def _from_blob(b: bytes):
    n = len(b) // 4
    return list(struct.unpack(f"<{n}f", b))


def _chunk(text: str):
    text = re.sub(r"\n{3,}", "\n\n", text.strip())
    return [text[i:i + CHUNK_SIZE] for i in range(0, len(text), CHUNK_SIZE)] or [""]


def embed_texts(cfg: dict, texts):
    """OpenAI 兼容批量嵌入（每批 ≤16 条，按 index 归位）。"""
    url = cfg["base_url"].rstrip("/")
    if not url.endswith("/embeddings"):
        url = url + "/embeddings"
    out = []
    for i in range(0, len(texts), 16):
        batch = texts[i:i + 16]
        r = _post_json(url, {"authorization": f"Bearer {cfg['api_key']}"},
                       {"model": cfg["model"], "input": batch})
        data = r.get("data", [])
        if len(data) != len(batch):
            raise RuntimeError(f"嵌入返回条数不符（期望 {len(batch)}，得 {len(data)}）")
        for it in sorted(data, key=lambda x: x.get("index", 0)):
            out.append(it["embedding"])
    return out


def embed_query(cfg: dict, q: str):
    return embed_texts(cfg, [q])[0]


@router.post("/embed")
def kb_embed(body: dict, u: dict = Depends(require_teacher)):
    """向量化：body.path 非空=单篇，否则全库（分块 → 调用嵌入 API → 落库）。"""
    cfg = _load_cfg("embed")
    if not cfg:
        raise HTTPException(400, "请先在「向量/重排模型」配置嵌入 API（地址/密钥/模型）")
    path = (body.get("path") or "").strip()
    docs = [(rel, text) for rel, text in llm_direct.kb_fresh(0) if (not path or rel == path)]
    if not docs:
        raise HTTPException(404, "目标文档不存在")
    texts, metas = [], []
    for rel, text in docs:
        for i, ch in enumerate(_chunk(text)):
            texts.append(ch)
            metas.append((rel, i, ch))
    try:
        vecs = embed_texts(cfg, texts)
    except HTTPException:
        raise
    except Exception as e:  # 网络/密钥/模型错误 → 明确报错，不静默
        raise HTTPException(502, f"嵌入 API 调用失败：{str(e)[:160]}")
    if len(vecs) != len(texts):
        raise HTTPException(502, f"嵌入返回条数不符（期望 {len(texts)}，得 {len(vecs)}）")
    now = int(time.time())
    d = db.get_db()
    try:
        for (rel, i, ch), v in zip(metas, vecs):
            d.execute(
                "INSERT INTO kb_embeddings(doc_path,chunk_idx,chunk_text,model,vector,created_at) "
                "VALUES(?,?,?,?,?,?) ON CONFLICT(doc_path,chunk_idx,model) DO UPDATE SET "
                "chunk_text=excluded.chunk_text, vector=excluded.vector, created_at=excluded.created_at",
                (rel, i, ch, cfg["model"], _to_blob(v), now))
        d.commit()
    finally:
        d.close()
    return {"ok": True, "docs": len(docs), "chunks": len(texts), "model": cfg["model"]}